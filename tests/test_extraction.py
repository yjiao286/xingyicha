#!/usr/bin/env python3
"""Unit tests for personnel / price extraction generalization.

Run: ./venv/bin/python3 tests/test_extraction.py
No pytest dependency; plain asserts with a tiny runner.
"""
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app as m  # noqa: E402

PASS = []


def check(label, fn):
    try:
        fn()
    except AssertionError as e:
        print(f'FAIL {label}: {e}')
        raise
    PASS.append(label)
    print(f'PASS {label}')


# ── Price extraction channels ──
def t_price_standalone_cn():
    r = m.extract_prices('人民币（大写）：壹亿贰仟陆佰壹拾捌万壹仟玖佰柒拾陆元叁角')
    assert r['totalPriceInTax'] is not None and abs(r['totalPriceInTax'] - 126181976.30) < 1, r


def t_price_symbol_with_wan():
    r = m.extract_prices('投标总价：￥12.5万元')
    assert abs(r['totalPriceInTax'] - 125000) < 1, r


def t_price_cn_label():
    r = m.extract_prices('投标总价：壹佰贰拾叁万元整')
    assert abs(r['totalPriceInTax'] - 1230000) < 1, r


def t_price_fullwidth():
    r = m.extract_prices('小写：１２３４５６')
    assert r['totalPriceInTax'] == 123456, r


def t_price_bond_excluded():
    r = m.extract_prices('投标保证金一份，金额为人民币：￥500000元')
    assert r['totalPriceInTax'] is None, r


def t_price_summary_sep():
    r = m.extract_prices('合计 \\ 1838529 5002800 \\')
    assert r['totalPrice'] == 1838529 and r['totalPriceInTax'] == 5002800, r


def t_price_bidrate():
    assert m.extract_prices('下浮率：12.5%')['bidRate'] == '12.5%'
    assert m.extract_prices('报价费率 1.8‰')['bidRate'] == '1.8‰'
    assert m.extract_prices('投标报价下浮6%')['bidRate'] == '6%'


def t_price_smallcase_first():
    r = m.extract_prices('小写：123456元 大写：壹拾贰万叁仟肆佰伍拾陆元整')
    assert r['totalPriceInTax'] == 123456, r


def t_price_cn_fractions():
    r = m.extract_prices('人民币（大写）：贰佰叁拾肆万伍仟陆佰柒拾捌元玖角')
    assert abs(r['totalPriceInTax'] - 2345678.9) < 0.01, r


def t_price_reverse_pair():
    r = m.extract_prices('小写：123456\n大写：壹拾贰万叁仟肆佰伍拾陆元整')
    assert r['totalPriceInTax'] == 123456, r


def t_similarity_toc_dots_ignored():
    # TOC leader dots ('四、授权委托书 ....... 7') must never pair unrelated
    # lines across documents ('9.7 网络拥塞的感知时间不高于 5s ....... 79')
    t1 = '目录\n四、授权委托书 ....... 7\n一、投标函 ......... 3\n'
    t2 = '9.7 网络拥塞的感知时间不高于 5s ....... 79\n项目人员与分工表\n'
    res = m.text_similarity_analysis({'a': t1, 'b': t2})
    for pr in res['pair_results']:
        for mm in pr.get('matches', []):
            assert '授权委托书' not in mm.get('text', ''), mm
            assert '网络拥塞' not in mm.get('text', ''), mm
    # 真实内容仍应匹配
    t3 = '本项目采用三层架构设计，核心交换节点采用双机热备冗余部署策略\n'
    t4 = '本项目采用三层架构设计，核心交换节点采用双机热备冗余部署策略\n'
    res2 = m.text_similarity_analysis({'a': t3, 'b': t4})
    assert any('三层架构' in mm.get('text', '')
               for pr in res2['pair_results'] for mm in pr.get('matches', [])), res2


def t_price_wan_near_context():
    # 万-form value must survive the near-context validation
    r = m.extract_prices('投标总价：￥12.5万元')
    assert abs(r['totalPriceInTax'] - 125000) < 1, r


def t_price_reference_amount_excluded():
    # 合同金额 reference amounts must never win over the real bid price
    t = ('开标一览表\n投标报价\n总价：7590000元\n\n业绩：\n'
         '项目名称：某某市计量质量检测研究院\n数量：1\n合同金额：RMB2080000.00\n')
    r = m.extract_prices(t)
    assert r['totalPriceInTax'] == 7590000, r
    # 招标文件售价 (document price, not bid price)
    r2 = m.extract_prices('开标一览表\n投标报价表\n\n招标文件售价：人民币1000元\n')
    assert r2['totalPriceInTax'] is None, r2


# ── Personnel extraction ──
def t_personnel_pipe_table():
    text = ('姓名 | 职务 | 电话\n'
            '张三 | 项目经理 | 13900000001\n'
            '李四 | 技术负责人 | 13900000002\n')
    p = m.extract_personnel(text)
    names = {x['name'] for x in p['all_persons']}
    assert {'张三', '李四'} <= names, names
    assert '13900000001' in p['phones'], p['phones']
    assert '13900000002' in p['phones'], p['phones']


def t_personnel_pipe_id():
    text = ('姓名 | 职务 | 身份证号\n王五 | 安全员 | 320101199001011234\n')
    p = m.extract_personnel(text)
    assert '320101199001011234' in p['id_numbers'], p['id_numbers']
    assert p['id_number'] == '320101199001011234', p['id_number']


def t_personnel_multivalue():
    p = m.extract_personnel('身份证号：320101199001011234 电话：13912345678 邮箱：a@b.com')
    assert '13912345678' in p['phones'] and '320101199001011234' in p['id_numbers']
    assert 'a@b.com' in p['emails'], p['emails']


def t_personnel_spaced_id():
    p = m.extract_personnel('身份证号：320101 1990 01 01 1234')
    assert p['id_numbers'] and p['id_numbers'][0] == '320101199001011234', p['id_numbers']


def t_personnel_auth_variants():
    p = m.extract_personnel('授权委托代理人：赵六\n联系电话：13800000000\n')
    assert p['authorized_rep'] == '赵六', p
    assert '13800000000' in p['phones'], p['phones']


def t_personnel_blank_template_no_leak():
    # Blank authorization-letter templates ('本人 （姓名）系 （投标人名称）…')
    # must not leak 本人/性别/盖单位章 as names or companies
    t = ('四、授权委托书\n本人 （姓名）系 （投标人名称）的法定代表人（单位负责人），现委托\n'
         '投标人：________________（盖单位章）\n姓名：性别：男 身份证号：\n')
    p = m.extract_personnel(t)
    names = {x['name'] for x in p['all_persons']}
    assert '本人' not in names and '性别' not in names, names
    assert p.get('legal_rep') not in ('本人', '性别', '性别:'), p.get('legal_rep')
    assert p.get('company_name') not in ('________________（盖单位章', '（盖单位章'), p.get('company_name')


def t_personnel_title_word_not_name():
    # PDF "序号 姓名 职称 分工" tables: '张伟 中级 项目负责人' — the 职称
    # column value 中级 must not become the name; 张伟 must be captured.
    t = ('项目人员配置\n表 5 项目人员与分工\n序号 姓名 职称 分工\n'
         '1 张伟 中级 项目负责人\n'
         '2 刘某某 教授 流资源预留协议设计\n'
         '3 潘某某 副教授 负载均衡技术设计\n')
    p = m.extract_personnel(t)
    names = [x['name'] for x in p['all_persons']]
    assert '中级' not in names and '教授' not in names and '副教授' not in names, names
    assert '张伟' in names, names
    zhang = [x for x in p['all_persons'] if x['name'] == '张伟']
    assert any(x['role'] == 'project_manager' for x in zhang), zhang


def t_personnel_role_not_name():
    # Role keywords in the name column must never become person entries
    p = m.extract_personnel('姓名 | 职务\n张三 | 项目经理\n项目经理 | 组长\n')
    names = {x['name'] for x in p['all_persons']}
    assert names == {'张三'}, names
    roles = {x['role'] for x in p['all_persons']}
    assert 'project_manager' in roles, roles


# ── xlsx ──
def t_xlsx_extract(tmp='/tmp/_t.xlsx'):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '报价单'
    ws.append(['序号', '名称', '数量', '单价', '总价'])
    ws.append([1, '服务器', 2, 30000, 60000])
    wb.save(tmp)
    try:
        text = m.extract_text_with_tables(tmp)
        assert '报价单' in text and '服务器' in text, text
        assert ' | ' in text
    finally:
        os.remove(tmp)


# ── Cross-match layer helpers (set-intersection pools) ──
def t_pools_intersection():
    from app import extract_personnel
    pi = extract_personnel('电话：13900000001 身份证号：320101199001011234')
    pj = extract_personnel('手机：13900000001 身份证号：320101199001011234')
    shared_phones = set(pi['phones']) & set(pj['phones'])
    shared_ids = set(pi['id_numbers']) & set(pj['id_numbers'])
    assert shared_phones == {'13900000001'}, shared_phones
    assert shared_ids == {'320101199001011234'}, shared_ids


# ── Format-generalization enhancements ──
def t_price_label_unit_paren():
    # '总价（元）：' and '合计（万元）：' label-with-unit formats
    r = m.extract_prices('开标一览表\n投标总价（元）：1230000元')
    assert r['totalPriceInTax'] == 1230000, r
    r2 = m.extract_prices('报价汇总表\n合计（万元）：89.3万元')
    assert abs(r2['totalPriceInTax'] - 893000) < 1, r2


def t_price_pipe_total_wan_cells():
    # 合计 pipe row with 万元-suffixed cells must keep magnitude
    r = m.extract_prices('序号 | 名称 | 金额\n合计 | 全部 | 89.3万元 | 91.97万元')
    assert abs(r['totalPrice'] - 893000) < 1, r
    assert abs(r['totalPriceInTax'] - 919700) < 1, r


def t_price_bidrate_extra():
    assert m.extract_prices('投标报价（%）：98.5')['bidRate'] == '98.5%', m.extract_prices('投标报价（%）：98.5')
    assert m.extract_prices('投标报价下浮 6 个百分点')['bidRate'] == '6%'


def t_personnel_zhweituo():
    p = m.extract_personnel('授权委托书\n兹委托 李勇 同志为我方代理人，负责签署投标文件。')
    assert p['authorized_rep'] == '李勇', p


def t_personnel_reversed_labels():
    p = m.extract_personnel('项目管理机构\n职务：项目经理 姓名：王强 联系电话：13800000000')
    names = {x['name'] for x in p['all_persons']}
    assert '王强' in names, names


def t_personnel_pipe_merged_cell():
    # docx vertically-merged name cells: empty name inherits row above
    text = ('姓名 | 职务 | 联系电话\n'
            '张三 | 项目经理 | 13900000001\n'
            ' | 技术负责人 | 13900000002\n')
    p = m.extract_personnel(text)
    names = {x['name'] for x in p['all_persons']}
    assert {'张三'} <= names, names
    assert '13900000002' in p['phones'], p['phones']


def t_personnel_pipe_multi_phone_cell():
    text = '姓名 | 职务 | 联系电话\n张三 | 项目经理 | 13900000001/13900000002\n'
    p = m.extract_personnel(text)
    assert '13900000001' in p['phones'] and '13900000002' in p['phones'], p['phones']


def t_personnel_pipe_header_alias():
    text = ('拟投入主要人员 | 职务\n'
            '李四 | 技术负责人\n')
    p = m.extract_personnel(text)
    names = {x['name'] for x in p['all_persons']}
    assert '李四' in names, names


def t_personnel_minority_dot_normalized():
    # Same minority name spelled with different middle dots must cross-match
    a = m.extract_personnel('项目管理机构\n姓名：阿不来提•买买提')
    b = m.extract_personnel('项目管理机构\n姓名：阿不来提·买买提')
    ka = {re.sub(r'[•・]', '·', x['name']) for x in a['all_persons']}
    kb = {re.sub(r'[•・]', '·', x['name']) for x in b['all_persons']}
    assert ka & kb, (ka, kb)


def t_personnel_bank_account_pool():
    p = m.extract_personnel('开户银行：中国工商银行北京分行\n账号：1100923456789000123')
    assert '1100923456789000123' in p.get('bank_accounts', []), p.get('bank_accounts')


def t_personnel_section_marker_extended():
    # New personnel-table section markers must scope extraction
    p = m.extract_personnel('人员一览表\n姓名 职务\n王强 项目经理\n')
    names = {x['name'] for x in p['all_persons']}
    assert '王强' in names, names


# ── Accuracy hardening: ceilings, environmental noise, cross-checks ──
def t_price_ceiling_excluded():
    # 最高限价 reprinted in every bid must never win as the bid price,
    # and must not HIDE the real price that appears further down.
    t = ('开标一览表\n最高限价（招标控制价）：￥1,000,000元\n投标报价：￥980,000元\n')
    r = m.extract_prices(t)
    assert r['totalPriceInTax'] == 980000, r


def t_price_ceiling_only_not_captured():
    r = m.extract_prices('开标一览表\n招标控制价：￥1000000元\n')
    assert r['totalPriceInTax'] is None, r


def t_bank_account_tender_side_excluded():
    # 保证金汇入账号 appears in EVERY bid → must stay out of the pool
    t = ('投标保证金请汇入以下账户\n开户银行：XX银行北京分行\n账号：1100923456789000123\n')
    p = m.extract_personnel(t)
    assert '1100923456789000123' not in p.get('bank_accounts', []), p.get('bank_accounts')


def t_environmental_pool_demotion():
    pools = {
        'A': {'phones': ['13900000001', '13800000000'], 'id_numbers': [], 'emails': [], 'bank_accounts': []},
        'B': {'phones': ['13900000002', '13800000000'], 'id_numbers': [], 'emails': [], 'bank_accounts': []},
        'C': {'phones': ['13900000003', '13800000000'], 'id_numbers': [], 'emails': [], 'bank_accounts': []},
        'D': {'phones': ['13900000004', '13800000000'], 'id_numbers': [], 'emails': [], 'bank_accounts': []},
        'E': {'phones': ['13800000001'], 'id_numbers': [], 'emails': [], 'bank_accounts': []},
    }
    removed = m._demote_environmental_pool_values(pools, list(pools))
    # 13800000000 in 4/5 groups (80%) → environmental, demoted
    assert '13800000000' in removed, removed
    assert pools['A']['phones'] == ['13900000001'], pools['A']
    assert pools['E']['phones'] == ['13800000001'], pools['E']
    # 3 of 4 (75%) is below the ratio — a genuinely shared phone between
    # 3 of 4 bidders is still collusion evidence and must NOT be demoted
    pools3 = {
        'A': {'phones': ['13900000001'], 'id_numbers': [], 'emails': [], 'bank_accounts': []},
        'B': {'phones': ['13900000001'], 'id_numbers': [], 'emails': [], 'bank_accounts': []},
        'C': {'phones': ['13900000001'], 'id_numbers': [], 'emails': [], 'bank_accounts': []},
        'D': {'phones': ['13900000009'], 'id_numbers': [], 'emails': [], 'bank_accounts': []},
    }
    assert not m._demote_environmental_pool_values(pools3, list(pools3))
    assert pools3['A']['phones'] == ['13900000001']
    # With 2 groups only, a shared phone stays (still real evidence)
    pools2 = {
        'A': {'phones': ['13900000001'], 'id_numbers': [], 'emails': [], 'bank_accounts': []},
        'B': {'phones': ['13900000001'], 'id_numbers': [], 'emails': [], 'bank_accounts': []},
    }
    assert not m._demote_environmental_pool_values(pools2, list(pools2))
    assert pools2['A']['phones'] == ['13900000001']


def t_price_cn_arabic_mismatch():
    # 大写 says 100万 but 小写 says 980000 → adopt 小写, note the conflict
    t = '开标一览表\n投标总价：壹佰万元整（小写：980000元）\n'
    r = m.extract_prices(t)
    assert r['totalPriceInTax'] == 980000, r
    assert any('大写' in w for w in r['warnings']), r['warnings']


def t_price_cn_arabic_consistent_no_warning():
    t = '开标一览表\n投标总价：壹佰万元整（小写：1000000元）\n'
    r = m.extract_prices(t)
    assert not any('大写' in w for w in r['warnings']), r['warnings']


def t_price_subitem_sum_warning():
    t = ('报价明细表\n序号 | 名称 | 数量 | 单价 | 总价\n'
         '1 | 服务器 | 2 | 30000 | 60000\n'
         '2 | 交换机 | 1 | 20000 | 20000\n'
         '合计 | | | | 980000\n')
    r = m.extract_prices(t)
    assert any('分项合计' in w for w in r['warnings']), r['warnings']


def t_price_global_fallback_warning():
    r = m.extract_prices('合计 \\ 1838529 5002800 \\')
    assert any('兜底' in w for w in r['warnings']), r['warnings']


def t_phone_spaced_and_ocr():
    p = m.extract_personnel('联系电话：139 1234 5678')
    assert '13912345678' in p['phones'], p['phones']
    p2 = m.extract_personnel('手机：138O1234567')
    assert '13801234567' in p2['phones'], p2['phones']


def t_surname_filter():
    # Non-surname table fragments must be rejected; rare/compound surnames kept
    assert not m._is_person_name('情况'), '情况 should be rejected'
    assert not m._is_person_name('规程'), '规程 should be rejected'
    assert m._is_person_name('欧阳建国')
    assert m._is_person_name('覃芳')
    assert m._is_person_name('诸葛云')


def t_parse_amount_ocr_letters():
    assert m._parse_amount('1O9,800') == 109800
    assert m._parse_amount('l2300') == 12300
    assert m._parse_amount('I00000') == 100000


def t_fitz_table_pipes():
    # PyMuPDF table rows → pipe lines (fake page object, no pymupdf needed)
    class FakeTable:
        def extract(self):
            return [['姓名', '职务'], ['张三', '项目经理'], [None, '']]
    class FakeFinder:
        tables = [FakeTable()]
    class FakePage:
        def find_tables(self):
            return FakeFinder()
    out = m._fitz_page_tables_as_pipes(FakePage())
    assert '姓名 | 职务' in out and '张三 | 项目经理' in out, out


# ── Real-world regression: auth-letter/cert format (2026-08-25) ──
def t_personnel_auth_letter_full():
    # 授权委托书 with labeled parens; also exercises the 兰 surname
    t = ('授权委托书\n本人  兰某某  （姓名）系  北京某某航天技术有限公司  （供应商名称）的'
         '法定代表人（单位负责人），现委托  陈某某  （姓名）为我方授权代理。\n')
    p = m.extract_personnel(t)
    assert p['legal_rep'] == '兰某某', p
    assert p['company_name'] == '北京某某航天技术有限公司', p
    assert p['authorized_rep'] == '陈某某', p


def t_personnel_cert_paren_format():
    # 法定代表人资格证明书: name AND company in unlabeled parentheses
    t = '法定代表人资格证明书\n（兰某某）系（北京某某航天技术有限公司）的法定代表人。\n特此证明\n'
    p = m.extract_personnel(t)
    assert p['legal_rep'] == '兰某某', p
    assert p['company_name'] == '北京某某航天技术有限公司', p


def t_personnel_company_survives_bad_name():
    # A rejected name must not also lose the company
    t = ('授权委托书\n本人  某某  （姓名）系  北京某某航天技术有限公司  （供应商名称）的法定代表人\n')
    p = m.extract_personnel(t)
    assert p['legal_rep'] is None, p
    assert p['company_name'] == '北京某某航天技术有限公司', p


def t_personnel_surnames_extended():
    for n in ('兰某某', '付某某', '肖某某', '闫某某', '郝某某', '滕某某', '岳某某', '单某某'):
        assert m._is_person_name(n), n


# ── History persistence: _prepare_history_data keep/strip semantics ──
def _fake_results():
    return {
        '_flags': {'group_map': {}, 'ok': True},
        'verdict': {'conclusion': '两份标书存在围串标高度嫌疑'},
        'metadata': {
            'files': [{'name': 'A.docx', 'creator': '张三',
                       'KSOProductBuildVer': '12.1.0.1', 'ICV': 'abc-def',
                       'KSOTemplateDocerSaveRecord': 'r1', 'pages': 12,
                       '_error': None, 'secret_huge': 'x' * 5000}],
            'matches': [], 'findings': [],
        },
        'personnel': {
            'files': [{'name': 'A.docx', 'legal_rep': '王五',
                       'phones': ['13800000000', '13911111111'],
                       'id_numbers': ['110101199001011234'],
                       'id_number': '110101199001011234',
                       'emails': ['a@b.com'], 'bank_accounts': ['110060123456'],
                       'all_persons': [{'name': '王五', 'role': 'legal_rep',
                                        'confidence': 0.9}],
                       'contacts': {'phone': '13800000000', 'email': None,
                                    'address': None},
                       'heavy_blob': ['x' * 3000]}],
            'cross_matches': [], 'findings': [],
        },
        'pricing': {
            'files': [{'name': 'A.docx', 'totalPriceInTax': 100,
                       'totalPrice': 92, 'taxRate': '13%', 'bidRate': '99.5%',
                       'revenue': 10, 'cost': 5,
                       'warnings': ['大写/小写不一致已修正'],
                       'subItemPrice': [{'priceName': '设备', 'totalPrice': 92}]}],
            'comparison': {},
            'subItemCompare': [
                {'name': '设备',
                 'items': [{'file': 'A.docx', 'totalPriceInTax': 100,
                            'extras': {'厂家/型号': '华为S5720'}}],
                 'findings': []}],
            'findings': [],
        },
        'text_similarity': {
            'pair_results': [{
                'file1': 'A', 'file2': 'B', 'total_matches': 1,
                'matches': [{'index': 1, 'length': 500, 'text': 'x' * 300,
                             'ctx1': 'y' * 500, 'ctx2': 'z' * 500,
                             'abnormal': True, 'risk_level': 'substantial',
                             'score': 0.8, 'reasons': ['术语密度高']}],
            }],
            'template_matches': 0, 'total_pairs': 1, 'findings': [],
        },
    }


def t_history_prep_keeps_new_fields():
    r = m._prepare_history_data(_fake_results())
    pf = r['personnel']['files'][0]
    assert pf['phones'] == ['13800000000', '13911111111'], pf
    assert pf['id_numbers'] == ['110101199001011234'], pf
    assert pf['emails'] == ['a@b.com'], pf
    assert pf['bank_accounts'] == ['110060123456'], pf
    assert pf['all_persons'][0]['name'] == '王五', pf
    qf = r['pricing']['files'][0]
    assert qf['bidRate'] == '99.5%', qf
    assert qf['revenue'] == 10 and qf['cost'] == 5, qf
    assert qf['warnings'] == ['大写/小写不一致已修正'], qf
    mf = r['metadata']['files'][0]
    assert mf['KSOProductBuildVer'] == '12.1.0.1', mf
    assert mf['KSOTemplateDocerSaveRecord'] == 'r1', mf
    assert mf['ICV'] == 'abc-def', mf


def t_history_prep_keeps_extras():
    r = m._prepare_history_data(_fake_results())
    items = r['pricing']['subItemCompare'][0]['items']
    assert items[0]['extras'] == {'厂家/型号': '华为S5720'}, items


def t_history_prep_light_matches():
    r = m._prepare_history_data(_fake_results())
    match = r['text_similarity']['pair_results'][0]['matches'][0]
    assert len(match['text']) == 200, len(match['text'])
    assert len(match['ctx1']) == 400 and len(match['ctx2']) == 400
    assert match['index'] == 1 and match['score'] == 0.8
    assert match['reasons'] == ['术语密度高']


def t_history_prep_strips_unknown_and_keeps_underscore():
    r = m._prepare_history_data(_fake_results())
    pf = r['personnel']['files'][0]
    # Unused by the frontend/report: contacts is a dict -> None
    assert pf['contacts'] is None, pf['contacts']
    assert pf['heavy_blob'] is None, pf['heavy_blob']
    # '_'-prefixed keys survive; unknown str keys become '' (type preserved)
    assert r['_flags'] == {'group_map': {}, 'ok': True}
    mf = r['metadata']['files'][0]
    assert mf['_error'] is None
    assert mf['secret_huge'] == '', repr(mf['secret_huge'])


def t_history_prep_modifies_copy_only():
    src = _fake_results()
    r = m._prepare_history_data(src)
    assert src['personnel']['files'][0]['phones'] == ['13800000000', '13911111111']
    assert src['text_similarity']['pair_results'][0]['matches'][0]['text'] == 'x' * 300
    assert src['pricing']['subItemCompare'][0]['items'][0]['extras'] is not None


# ── Company-name / authorized-rep cleanliness regressions ──
def t_company_strips_detached_label():
    # '（投标人名称' had the closing paren consumed by a wider capture — the
    # detached fragment must be cleaned away, not left on the company name.
    assert m._clean_company('北京某某大学 （投标人名称') == '北京某某大学'
    assert m._clean_company('北京某某大学（盖单位章）') == '北京某某大学'
    assert m._clean_company('北京某某大学（盖单位章') == '北京某某大学'
    assert m._clean_company('___北京某某大学___（盖单位章）') == '北京某某大学'
    assert m._clean_company('北京某某航天技术有限公司') == '北京某某航天技术有限公司'


def t_auth_xianweituo_split_newline():
    # '现委托李四为我方代理\n人' — 代理人 split across a PDF line wrap must
    # still be recognized (pattern 8 tolerates the newline).
    info = {'legal_rep': None, 'authorized_rep': None, 'company_name': None, 'all_persons': []}
    m._extract_from_auth_section(
        '本人张某某系北京某某大学的法定代表人（单位负责人），现委托李四为我方代理\n人。'
        '代理人根据授权，以我方名义签署、澄明确认。', info)
    assert info['authorized_rep'] == '李四', info
    assert info['legal_rep'] == '张某某', info
    assert info['company_name'] == '北京某某大学', info


def t_history_prep_keeps_name():
    r = m._prepare_history_data(_fake_results())
    assert r['personnel']['files'][0]['name'] == 'A.docx', r['personnel']['files'][0]['name']
    assert r['pricing']['files'][0]['name'] == 'A.docx'
    assert r['metadata']['files'][0]['name'] == 'A.docx'


def _glue(text):
    return m._glue_phrases(text)


def t_glue_phrases_multi_position():
    # 法定\n代\n表\n人 split across several line breaks -> one keyword.
    assert _glue('本人张某某系北京某\n某大学的法定\n代\n表\n人（单位负责人）') == \
        '本人张某某系北京某\n某大学的法定代表人（单位负责人）'
    # 委托代\n理人
    assert _glue('现委托李四为我方委托代\n理人。代理人行使签署权。') == \
        '现委托李四为我方委托代理人。代理人行使签署权。'
    # absent keyword is untouched
    assert _glue('这是一段普通文字，没有关键词') == '这是一段普通文字，没有关键词'


def t_glue_cleans_auth_extraction():
    # A whole-keyword word wrap inside the auth letter must not lose the agent.
    info = {'legal_rep': None, 'authorized_rep': None, 'company_name': None, 'all_persons': []}
    m._extract_from_auth_section(
        '本人张某某系北京某某大学的法定代表人（单位负责人），现委托李四为我方代理\n人。', info)
    assert info['authorized_rep'] == '李四', info
    assert info['company_name'] == '北京某某大学', info


def t_cjk_ws_normalized():
    assert m._normalize_cjk_whitespace('投标人 ： 张三（ 盖单位章 ）') == '投标人： 张三（盖单位章）'
    assert m._normalize_cjk_whitespace('电话 ：　010-12345678') == '电话： 010-12345678'


def t_join_split_names():
    # 换行拆词拼接
    assert m._join_split_names('本人张\n某某系的法人') == '本人张某某系的法人'
    # 空格列间距不得拼接（否则 '国 联系' / '建国 联系' 吞掉标签）
    assert m._join_split_names('张 某某') == '张 某某'
    assert m._join_split_names('职务：项目经理 姓名：王强 联系电话：13800000000') == \
        '职务：项目经理 姓名：王强 联系电话：13800000000'


def t_glue_new_phrases():
    assert m._glue_phrases('投标\n文件：开标一览表') == '投标文件：开标一览表'
    assert m._glue_phrases('授权\n委托\n书') == '授权委托书'


def t_company_prefix_strip():
    assert m._clean_company('投标人：北京某某大学') == '北京某某大学'
    assert m._clean_company('单位名称：北京某某大学') == '北京某某大学'
    assert m._clean_company('企业名称　某省未来网络创新研究院') == '某省未来网络创新研究院'


def t_clean_phone_junk():
    assert m._clean_phone('_021-12345678、') == '021-12345678'
    assert m._clean_phone('010-12345678;') == '010-12345678'
    assert m._clean_phone('010 - 1234 5678') == '010-12345678'


# ── Defense layers 7+: invisible chars / OCR glyphs / phone & amount forms ──
def t_invisible_chars_stripped():
    # Zero-width chars sit inside labels ('委托代理人\u200b：'), where \s-based
    # glue cannot reach them; CR/form-feed variants must become newlines.
    t = '授权委托书\r\n委托代理人\u200b：李\u2060勇\r被授权人：王\u00ad强\n'
    p = m.extract_personnel(t)
    assert p['authorized_rep'] == '李勇', p


def t_ocr_label_glyph_fixes():
    # Scanned-label glyph confusions: 人/入, 话/活, 币/巾, plus the 身分证
    # variant spelling — all repaired before any label regex runs.
    p = m.extract_personnel('授权委托书\n法定代表入：王强\n联系电活：13912345678\n'
                            '身分证号：110101199003070000\n')
    assert p['legal_rep'] == '王强', p
    assert '13912345678' in p['phones'], p
    assert '110101199003070000' in p['id_numbers'], p
    r = m.extract_prices('人民巾（大写）：壹佰贰拾万元整')
    assert r['totalPriceInTax'] == 1200000, r


def t_phone_country_code_and_dashed():
    # '+86…' is rejected by the plain pattern's (?<!\d) lookbehind; the
    # dash-grouped form never matched any pattern before.
    p = m.extract_personnel('联系人：+8613912345678 或 139-1234-5678')
    assert '13912345678' in p['phones'], p['phones']
    assert p['phones'].count('13912345678') == 1, p['phones']


def t_phone_landline_padded_and_label_variants():
    p = m.extract_personnel('授权委托书\n联系电话：010 - 1234 5678\n')
    assert p['phone'] == '010-12345678', p
    p2 = m.extract_personnel('授权委托书\n移动电话：13912345678\n')
    assert p2['phone'] == '13912345678', p2


def t_price_thinspace_groups():
    r = m.extract_prices('投标总价：￥1 261 819.76元')
    assert abs(r['totalPriceInTax'] - 1261819.76) < 0.01, r
    # Two space-separated column numbers must NOT merge into one amount.
    assert m._parse_amount('1838529 5002800') == 1838529


def t_price_thinspace_survives_validation():
    # The full chain: section detection must accept space-grouped numbers,
    # and the post-validation "value must appear in text" check must search
    # the '1 234 567' form — otherwise the extracted price is cleared again.
    r = m.extract_prices('开标一览表\n投标总价（元）：1 234 567.89\n税率 6%')
    assert r['totalPriceInTax'] == 1234567.89, r


def t_name_internal_spaces():
    p = m.extract_personnel('授权委托书\n委托代理人：张 三\n')
    assert p['authorized_rep'] == '张三', p
    p2 = m.extract_personnel('项目管理机构\n姓名：王 强 联系电话：13800000000\n')
    names = {x['name'] for x in p2['all_persons']}
    assert '王强' in names and '王强联' not in names, names


def t_name_tolerant_no_label_swallow():
    # '兹委托 李勇 同志…' — the tolerant class must not absorb 同志 even
    # though the trailing (?:同志)? is optional; the spaced form must still
    # be captured.
    p = m.extract_personnel('授权委托书\n兹委托 李勇 同志为我方代理人，负责签署投标文件。')
    assert p['authorized_rep'] == '李勇', p
    p2 = m.extract_personnel('授权委托书\n兹委托 李 勇 同志为我方代理人。')
    assert p2['authorized_rep'] == '李勇', p2


def t_name_validation_allows_internal_space():
    assert m._is_person_name('张 三')
    assert m._is_person_name('阿不来提 · 买买提')


def t_address_cut_at_next_label():
    p = m.extract_personnel('授权委托书\n地址：北京市海淀区上地十街10号 电话：010-12345678\n')
    assert p['address'] == '北京市海淀区上地十街10号', p


def t_email_fullwidth_at():
    p = m.extract_personnel('电子邮箱：Zhang＠Example.com\n')
    assert 'zhang@example.com' in p['emails'], p['emails']


def t_duty_label_suffix_not_person():
    # Duty-label strings are form/table column labels, never person names,
    # even when they start with a real surname (任中职务 → 任).
    for lbl in ('任中职务', '现任职务', '担任职务', '任何职务', '拟任职务',
                '本人岗位', '项目职称', '担任角色'):
        assert not m._is_person_name(lbl), lbl
    # Real names sharing those surnames/characters must keep passing.
    assert m._is_person_name('任志强')
    assert m._is_person_name('陈刚')


def t_personnel_duty_label_column():
    # PDF-flattened row '陈刚 任中职务 项目经理' (colons lost, label column
    # between name and role): the label must not enter all_persons, and the
    # real name before it must be recovered as project_manager.
    p = m.extract_personnel('项目团队\n陈刚 任中职务 项目经理\n杨帆 任中职务 项目经理\n')
    names = [(x['name'], x['role']) for x in p['all_persons']]
    assert ('任中职务', 'project_manager') not in names, names
    assert ('陈刚', 'project_manager') in names, names
    assert ('杨帆', 'project_manager') in names, names
    # Cross-line re-pairing must not duplicate an exact (name, role) pair.
    assert names.count(('杨帆', 'project_manager')) == 1, names


def t_personnel_duty_label_colon_form():
    p = m.extract_personnel('项目团队\n姓名：陈刚 任中职务：项目经理\n')
    names = [(x['name'], x['role']) for x in p['all_persons']]
    assert ('陈刚', 'project_manager') in names, names
    assert all(x['name'] != '任中职务' for x in p['all_persons']), names


def t_name_rejects_section_header_words():
    # Section markers / table captions that START with real surnames
    # (项/关/管) must never validate as person names.
    for hdr in ('项目团队', '项目成员', '项目人员', '关键人员', '管理人员',
                '项目分工', '项目组', '项目部', '组织机构', '人员简历',
                '项目机构', '人员配置'):
        assert not m._is_person_name(hdr), hdr
    # Real names sharing those surnames must keep passing.
    assert m._is_person_name('管虎')
    assert m._is_person_name('项少龙')


def t_name_rejects_duty_label_variants():
    # Traditional-script duty labels survive the surname check when led by
    # a surname char ('任職務') — the suffix rule must catch them too.
    for lbl in ('任職務', '任何崗位', '现任职务', '拟任职务', '岗位职责',
                '本任職責', '擔任角色'):
        assert not m._is_person_name(lbl), lbl


def t_personnel_section_header_not_person():
    # Flattened '项目团队\n项目经理：赵四' must not yield the section header
    # 项目团队 as a project_manager (项 is a real surname).
    p = m.extract_personnel('项目团队\n项目经理：赵四\n')
    names = [(x['name'], x['role']) for x in p['all_persons']]
    assert ('项目团队', 'project_manager') not in names, names
    assert ('赵四', 'project_manager') in names, names


def t_personnel_role_not_paired_across_lines():
    # A role keyword at the end of one table row must not attach to the
    # next row's name ('王强 项目经理\n李勇 施工员' once made 李勇 a
    # project_manager → false same-role cross-file match risk).
    p = m.extract_personnel('项目团队\n王强 项目经理\n李勇 施工员\n')
    roles = {}
    for x in p['all_persons']:
        roles.setdefault(x['name'], set()).add(x['role'])
    assert 'project_manager' in roles.get('王强', set()), roles
    assert 'project_manager' not in roles.get('李勇', set()), roles
    assert 'team_member' in roles.get('李勇', set()), roles
    # Same-line role-then-name pairs keep working, several per line.
    p2 = m.extract_personnel('项目团队\n项目经理 王强 技术负责人 李四\n')
    roles2 = {}
    for x in p2['all_persons']:
        roles2.setdefault(x['name'], set()).add(x['role'])
    assert roles2.get('王强') == {'project_manager'}, roles2
    assert roles2.get('李四') == {'tech_lead'}, roles2


def t_personnel_label_colon_without_name_label():
    # '陈刚 任中职务：项目经理' — no 姓名 label at all; the colon sits after
    # the duty-label column with zero space before the role value.
    p = m.extract_personnel('项目团队\n陈刚 任中职务：项目经理\n')
    names = [(x['name'], x['role']) for x in p['all_persons']]
    assert ('陈刚', 'project_manager') in names, names
    assert all(x['name'] != '任中职务' for x in p['all_persons']), names


def t_personnel_traditional_duty_label():
    # Traditional-script duty labels (任職務) must be skipped the same way
    # so the real name before them is still recovered.
    p = m.extract_personnel('项目团队\n陈刚 任職務 项目经理\n')
    names = [(x['name'], x['role']) for x in p['all_persons']]
    assert ('陈刚', 'project_manager') in names, names
    assert all('職務' not in x['name'] for x in p['all_persons']), names


def t_personnel_resume_vertical_labels():
    # Resume forms put 姓名 and 职务 on different lines; the span between
    # them may cross a newline but never another 姓名 label.
    p = m.extract_personnel('项目团队\n姓名：张三 性别：男\n现任职务：项目经理\n')
    roles = {x['role'] for x in p['all_persons'] if x['name'] == '张三'}
    assert 'project_manager' in roles, p['all_persons']

    p2 = m.extract_personnel('项目团队\n姓名：张三 学历：本科\n姓名：李四 职务：项目经理\n')
    roles2 = {x['role'] for x in p2['all_persons'] if x['name'] == '张三'}
    roles3 = {x['role'] for x in p2['all_persons'] if x['name'] == '李四'}
    assert 'project_manager' not in roles2, p2['all_persons']
    assert 'project_manager' in roles3, p2['all_persons']


def t_personnel_fill_brackets_unwrapped():
    # 【】-wrapped form fills ('姓名： 【张三】') must not defeat the label
    # regexes — one bidder wrapped EVERY filled value that way and personnel
    # came back completely empty.
    t = ('二、法定代表人（单位负责人）身份证明\n'
         '投标人名称： 某某机电制造有限公司\n'
         '姓名： 【张三】 性别： 【男】 年龄： 【76】 职务： 【总经理】\n'
         '\n'
         '3\n'
         '系 某某机电制造有限公司 （投标人名称）的法定代表人（单位\n'
         '负责人）。\n'
         '特此证明。\n')
    r = m.extract_personnel(t)
    assert r['company_name'] == '某某机电制造有限公司', r
    assert r['legal_rep'] == '张三', r


def t_person_spaced_authorized_rep_with_ocr_heading():
    # Scanned-bid regression: the authorization letter heading OCR'd to
    # '售权委托书' and the agent name was printed with per-char spaces
    # ('现委托 王 小 明 为我方代理人') — the agent was silently dropped, so
    # the cross-bid match against the same clean-spelled name never fired.
    t = ('本人 李林\n'
         '三、售权委托书\n'
         '（姓名） 系 某某管路配件有限公司 （投标人名称） 的法定代表人（单\n'
         ' 位负责人） ，现委托 王 小 明 为我方代理人。代理人根据授权，以我方名义签署、澄清\n'
         '确认、递交、撤回、修改设备采购招标项目投标文件、签订合同和处理有关事宜，其法律后\n'
         '果由我方承担。\n'
         '投标人：某某管路配件有限公司（盖单位章）\n')
    r = m.extract_personnel(t)
    assert r['company_name'] == '某某管路配件有限公司', r
    assert r['legal_rep'] == '李林', r
    # Spaced capture must be collapsed so cross-file name matching pairs it
    # with the clean spelling used in the other bidder's document.
    assert r['authorized_rep'] == '王小明', r


def t_person_auth_header_interleave_not_midword():
    # When the chapter header sits between the name line and the （姓名）
    # continuation, the capture must be the real name — never a mid-word
    # slice of the heading like '权委托书' (out of '三、授权委托书').
    t = ('本人 李林\n'
         '三、授权委托书\n'
         '（姓名） 系 某某管路配件有限公司 （投标人名称） 的法定代表人（单\n'
         ' 位负责人）。\n')
    r = m.extract_personnel(t)
    assert r['legal_rep'] == '李林', r


def t_person_join_split_names_spares_chapter_marker():
    # '李林\n三、授权委托书' is a name line followed by a chapter header —
    # joining yields '李林三、…' and derails the auth patterns. Numeral
    # continuations directly followed by a pause mark must not join.
    s = '本人 李林 \n三、授权委托书\n（姓名） 系 某某公司 的法定代表人'
    assert m._join_split_names(s) == s, m._join_split_names(s)
    # Genuine split names still join.
    assert m._join_split_names('负责人：张\n三丰') == '负责人：张三丰'


def t_ocr_heading_shouquan_fixed():
    # 授/售 glyph confusion in the authorization heading ('售权委托' never
    # occurs in legitimate Chinese) — fixing it restores section detection.
    assert m._fix_ocr_label_confusions('三、售权委托书') == '三、授权委托书'
    assert '授权' in m._fix_ocr_label_confusions('售权委托书')


def t_price_fill_brackets_unwrapped():
    # Same fill style hiding amounts: '（￥ 【98000】 ）' / 大写 in 【】.
    t = ('开标一览表\n投标总价（大写）： 【玖万捌仟元整】\n'
         '小写：￥ 【98000】 \n')
    r = m.extract_prices(t)
    assert r['totalPriceInTax'] == 98000, r


def t_person_xlsx_sheet_marker_survives_bracket_strip():
    # The 【工作表】 header drives the pipe-table parser — the bracket strip
    # must not eat it.
    s = '【工作表】人员表\n姓名 | 电话\n张三 | 13912345678\n'
    assert m._strip_fill_brackets(s) == s, m._strip_fill_brackets(s)


def t_price_bank_voucher_excluded():
    # Bond transfer slip: '金额 / 人民币：22,000.00 / 人民币：贰万贰仟元整 /
    # 用途 / 制单日期…' — banking words sit on neighbouring lines, not on
    # the amount's own line, so they need the voucher WINDOW check.
    text = ('开标一览表\n投标总价\n大写：壹佰壹拾贰万元\n小写：1120000.00元\n（元，含税）\n'
            '账号 30210000000123\n开户行 中国民生银行\n金额\n人民币：22,000.00\n'
            '人民币：贰万贰仟元整\n用途\n制单日期：2025-08-26\n')
    r = m.extract_prices(text)
    assert r['totalPriceInTax'] == 1120000, r


def t_price_bond_echo_excluded():
    # The bond value echoes unlabeled later in the document; a candidate
    # equal to a bond-labeled amount must be skipped even without voucher
    # words around it.
    text = '投标函\n投标总价：980000元\n投标保证金\n2.2万元\n附证明材料\n人民币：22000\n'
    r = m.extract_prices(text)
    assert r['totalPriceInTax'] == 980000, r


def t_price_rmb_label_no_newline_grab():
    # '2000万元人民币\n2014年1月20日' — the 人民币 prefix at a line end must
    # not pair with the next line's year; '人民币 2014年' (same line) is a
    # year too, never an amount.
    text = ('投标人基本情况表\n注册资金\n成立时间\n2000万元人民币\n2014年1月20日\n'
            '人民币 2014年注册\n开标一览表\n小写：1120000.00元\n')
    r = m.extract_prices(text)
    assert r['totalPriceInTax'] == 1120000, r


def t_price_section_toc_and_quote_anchors():
    # A TOC without dot leaders, a 《…》 reference and a "…" quoted mention
    # are not section headers; the anchor must fall on the real table.
    toc = '目录\n' + ''.join(
        f'{c}、{t}\n' for c, t in zip('一二三四五六七八九十',
        ['投标函', '开标一览表', '分项报价表', '法定代表人身份证明', '资格审查资料',
         '技术方案', '服务承诺', '履约计划', '培训方案', '其他补充材料']))
    toc += '附件1 本项目实施团队主要人员名单 63\n附件2 本项目实施团队主要人员简历表 64\n' * 6
    text = (toc + '一、投标函\n愿意以《开标一览表》中的投标报价提供全部内容。\n'
            '备注：分项合计金额中的总价应等于“开标一览表”中的投标总价。\n'
            '开标一览表\n投标总价（元，含税）\n大写：玖拾捌万元整\n小写：980000.00元\n')
    r = m.extract_prices(text)
    assert r['totalPriceInTax'] == 980000, r


def t_price_cost_wan_magnitude():
    r = m.extract_prices('成本明细\n人工成本 350万元\n设备购置费 12000元\n')
    costs = {d['priceName']: d['totalPrice'] for d in r['costDetails']}
    assert costs.get('人工成本') == 3500000, costs
    assert costs.get('设备购置费') == 12000, costs


def t_price_cost_performance_table_excluded():
    # Past-performance tables list HISTORICAL contract amounts under a
    # 合同金额 header; project-name fragments ('站产品采购') are not cost
    # items and 3500万 must not shrink to 3500.
    text = ('业绩一览表\n项目名称\n合同金额\n(元)\n2023年\n某某A+福利\n'
            '站产品采购\n3500万\n中国移动\n2000万\n')
    r = m.extract_prices(text)
    assert r['cost'] is None and not r['costDetails'], r


def t_price_unit_rate_not_total():
    # Per-month / per-person rates share the summary table with the total;
    # the denominator is dropped by the amount patterns, so the rate would
    # pose as the total.
    text = ('开标一览表\n小写：22000.00元/月\n服务单价 340000/人月\n'
            '投标总价：1120000元\n')
    r = m.extract_prices(text)
    assert r['totalPriceInTax'] == 1120000, r


def t_price_fee_penalty_labels_excluded():
    text = ('投标函\n违约金：人民币50000元\n招标代理服务费：￥30000\n'
            '开标一览表\n小写：980000.00元\n')
    r = m.extract_prices(text)
    assert r['totalPriceInTax'] == 980000, r


def t_price_bond_echo_cn_numeral():
    # '投标保证金（大写）：贰万元整' must seed the bond echo set so the
    # unlabeled '人民币：20000' echo is skipped anywhere in the document.
    text = ('投标保证金（大写）：贰万元整\n其他材料\n人民币：20000\n'
            '开标一览表\n小写：980000.00元\n')
    r = m.extract_prices(text)
    assert r['totalPriceInTax'] == 980000, r


def t_price_tax_pair_date_runs_rejected():
    # '签订 20250826 / 生效 20250829 / 3' — two calendar dates + a small
    # digit pass the old ratio check (1.0) and would pose as a 20M price pair.
    r = m.extract_prices('开标一览表\n签订 20250826\n生效 20250829\n3\n')
    assert r['totalPriceInTax'] is None and r['totalPrice'] is None, r


def t_price_cost_line_flattened_variants():
    # 序号 prefixes, colons, thousands commas and 万元 magnitudes are all
    # common in flattened cost tables.
    r = m.extract_prices('成本明细\n1. 材料费：340,000.00元\n（二）人工成本：350万\n')
    costs = {d['priceName']: d['totalPrice'] for d in r['costDetails']}
    assert costs.get('材料费') == 340000, costs
    assert costs.get('人工成本') == 3500000, costs


def t_price_cost_date_run_rejected():
    # Glued date runs ('检测20250826批次') are not 20M cost items.
    r = m.extract_prices('检测报告\n检测20250826批次\n差旅费20250826\n')
    assert not r['costDetails'] and r['cost'] is None, r


# ══ 第三批泛化：人员（填空下划线 / 无冒号标签 / 动词族 / 联系人 / 跨行断裂）══

def t_personnel_underscore_fill_names():
    # Values written over an underline fill survive label matching and are
    # stripped in the result.
    p = m.extract_personnel('授权委托书\n法定代表人：___张三___\n委托代理人：____李四____\n')
    assert p['legal_rep'] == '张三', p
    assert p['authorized_rep'] == '李四', p


def t_personnel_legal_rep_paren_sign():
    # '法定代表人（签字）：王五' — paren qualifier between label and colon.
    p = m.extract_personnel('投标函\n法定代表人（签字）：王五\n')
    assert p['legal_rep'] == '王五', p


def t_personnel_signature_deputy_variant():
    # '法定代表人或委托代理人' (missing 其) + （签字） qualifier + underline fill.
    p = m.extract_personnel('签字盖章页\n法定代表人或委托代理人（签字）：___赵六___\n')
    names = {x['name'] for x in p['all_persons']}
    assert '赵六' in names, names


def t_personnel_weipai_verb():
    # 委派/委任 verb variants; lazy capture keeps 同志 out of the name.
    p = m.extract_personnel('授权委托书\n兹委派孙七同志为我方代理人\n')
    assert p['authorized_rep'] == '孙七', p
    p2 = m.extract_personnel('授权委托书\n兹委任王五为我方代理人\n')
    assert p2['authorized_rep'] == '王五', p2


def t_personnel_contact_person_role():
    # '联系人：周八' feeds all_persons as bid_contact (cross-file same-name
    # matching previously never saw contact persons).
    p = m.extract_personnel('投标函\n联系人：周八\n联系电话：13912345678\n')
    roles = {x['name']: x['role'] for x in p['all_persons']}
    assert roles.get('周八') == 'bid_contact', roles


def t_personnel_colonless_legal_rep():
    # Flattened-table form '法定代表人 王大锤' (no colon).
    p = m.extract_personnel('兹委任某人\n法定代表人 王大锤\n')
    assert p['legal_rep'] == '王大锤', p


def t_personnel_id_newline_split():
    # Cell-per-line PDF splits an 18-digit ID across three lines.
    p = m.extract_personnel('身份证号：320123\n19900101\n123X\n')
    assert '32012319900101123X' in p['id_numbers'], p['id_numbers']


def t_personnel_mobile_group_linebreak():
    # '1391234\n5678' — 3-4-4 grouping split at a line break.
    p = m.extract_personnel('联系电话：1391234\n5678\n')
    assert '13912345678' in p['phones'], p['phones']


def t_personnel_role_vocab_extended():
    # Supervision/construction roles: 监理工程师 / 施工负责人.
    p = m.extract_personnel('主要人员\n李芳 监理工程师\n王雷 施工负责人\n')
    roles = {}
    for x in p['all_persons']:
        roles.setdefault(x['name'], set()).add(x['role'])
    assert 'team_member' in roles.get('李芳', set()), roles
    assert 'project_manager' in roles.get('王雷', set()), roles


def t_personnel_dunhao_name_role():
    # '张三、项目经理' — 顿号 separator between name and role.
    p = m.extract_personnel('项目团队\n张三、项目经理\n')
    roles = {x['name']: x['role'] for x in p['all_persons']}
    assert roles.get('张三') == 'project_manager', roles


def t_personnel_company_label_with_name():
    # Cover company label variant '投标人名称：…'.
    p = m.extract_personnel('投标人名称：北京某某科技有限公司\n')
    assert p['company_name'] == '北京某某科技有限公司', p.get('company_name')


def t_personnel_address_fill_stripped():
    p = m.extract_personnel('授权委托书\n地址：____北京市海淀区中关村大街1号____\n')
    assert p['address'] == '北京市海淀区中关村大街1号', p.get('address')


def t_personnel_bank_card_label():
    p = m.extract_personnel('银行卡号：6222020200112233445\n')
    assert '6222020200112233445' in p['bank_accounts'], p['bank_accounts']


def t_personnel_colonless_name_role_table():
    # '姓名 张三 职务 项目经理' — colon-less flattened table row.
    p = m.extract_personnel('项目团队\n姓名 张三 职务 项目经理\n')
    roles = {}
    for x in p['all_persons']:
        roles.setdefault(x['name'], set()).add(x['role'])
    assert 'project_manager' in roles.get('张三', set()), roles


# ══ 第三批泛化：报价（填空 / 括号变体 / 标签同义词 / 大写跨行 / 千分位）══

def t_price_underscore_fill():
    r = m.extract_prices('开标一览表\n投标总价：____98.6万元____\n')
    assert r['totalPriceInTax'] == 986000, r


def t_price_label_paren_note():
    # Short annotation paren between label and colon.
    r = m.extract_prices('开标一览表\n投标总价（含税）：1,261,819.76元\n')
    assert abs((r['totalPriceInTax'] or 0) - 1261819.76) < 0.01, r


def t_price_paren_unit_rmb():
    r = m.extract_prices('开标一览表\n总报价（人民币）：500000元\n')
    assert r['totalPriceInTax'] == 500000, r


def t_price_paren_unit_wan_bare_value():
    # Unit lives ONLY in the paren — bare 89.3 must scale to 893000.
    r = m.extract_prices('开标一览表\n合计（单位：万元）：89.3\n')
    assert r['totalPriceInTax'] == 893000, r
    r2 = m.extract_prices('报价汇总表\n合计（万元）：89.3\n')
    assert r2['totalPriceInTax'] == 893000, r2


def t_price_label_synonyms():
    for text, want in [('磋商报价：1200000元\n', 1200000),
                       ('最后报价：98万元\n', 980000),
                       ('报价总额：350万\n', 3500000)]:
        r = m.extract_prices(text)
        assert r['totalPriceInTax'] == want, (text, r)


def t_price_daxie_split_across_lines():
    # CN amount broken mid-number by a line break.
    r = m.extract_prices('大写：壹佰贰拾\n叁万元整\n')
    assert r['totalPriceInTax'] == 1230000, r


def t_price_daxie_jine_word():
    # '大写金额：' inserts 金额 between label and colon.
    r = m.extract_prices('大写金额：壹佰万元整\n')
    assert r['totalPriceInTax'] == 1000000, r


def t_price_daxie_split_pair_consistent():
    # Split 大写 + 小写 pair: value correct, no mismatch warning.
    r = m.extract_prices('大写：壹佰贰拾\n叁万元整\n小写：1230000元\n')
    assert r['totalPriceInTax'] == 1230000, r
    assert not any('不一致' in w for w in r['warnings']), r['warnings']


def t_price_xiaoxie_class_separator():
    # '（小写）￥：980000元' — ￥ before the colon, class-run separator.
    r = m.extract_prices('（小写）￥：980000元\n')
    assert r['totalPriceInTax'] == 980000, r


def t_price_heji_row_commas():
    # 合计 row with thousands commas in both price columns.
    r = m.extract_prices('合计 \\ 1,838,529 5,002,800 \\\n')
    assert r['totalPrice'] == 1838529 and r['totalPriceInTax'] == 5002800, r


def t_price_tax_included_wan_magnitude():
    r = m.extract_prices('含税总价：98万元\n')
    assert r['totalPriceInTax'] == 980000, r


def t_price_tax_decomposition_commas():
    # Pre-tax / post-tax pair with comma thousands: 1,838,529 × 1.03.
    r = m.extract_prices('开标一览表\n1,838,529 1893684 3\n')
    assert r['totalPrice'] == 1838529, r
    assert r['totalPriceInTax'] == 1893684, r
    assert r['taxRate'] == '3%', r


def t_price_rmb_lowercase_token():
    # 'rmb 98000元' — lowercase currency token.
    r = m.extract_prices('rmb 98000元\n')
    assert r['totalPriceInTax'] == 98000, r


def t_price_bidrate_underscore_fill():
    r = m.extract_prices('报价函\n下浮率：__8__%\n')
    assert r['bidRate'] == '8%', r


def t_price_xiaoxie_date_rejected():
    # A date run after 小写： must not become a 20M price.
    r = m.extract_prices('开标记录表\n小写：20250826\n')
    assert r['totalPriceInTax'] is None, r


# ══ 第四批泛化：括号注记组合 / 传真 / 为是连接 / 繁体大写 / 费率变体 ══

def t_personnel_paren_combined_qualifier():
    # '（签字或盖章）' combined paren qualifier on the agent label.
    p = m.extract_personnel('投标函\n委托代理人（签字或盖章）：钱九\n')
    assert p['authorized_rep'] == '钱九', p
    # '法定代表人（单位负责人）：' insert-paren label variant.
    p2 = m.extract_personnel('投标函\n法定代表人（单位负责人）：孙十\n')
    assert p2['legal_rep'] == '孙十', p2


def t_personnel_fax_paren_area_code():
    # 传真 label + '(010)1234-5678' parenthesized area code (parens dropped,
    # dash structure kept per _clean_phone convention).
    p = m.extract_personnel('投标函\n传真：(010)1234-5678\n')
    assert p['phone'] == '0101234-5678', p.get('phone')


def t_personnel_bare_zhanghu_label():
    p = m.extract_personnel('授权委托书\n账户：1100923456789001\n')
    assert '1100923456789001' in p['bank_accounts'], p['bank_accounts']


def t_personnel_cover_address():
    # Address on the bid-letter cover (previously only auth sections scanned).
    p = m.extract_personnel('投标函\n地址：北京市朝阳区建国路88号院6号楼\n')
    assert p['address'] == '北京市朝阳区建国路88号院6号楼', p.get('address')


def t_personnel_bare_auth_titles():
    p = m.extract_personnel('授权书\n兹授权周明为我方代理人\n')
    assert p['authorized_rep'] == '周明', p


def t_personnel_contact_fill():
    p = m.extract_personnel('投标函\n联系人：___郑一___\n')
    roles = {x['name']: x['role'] for x in p['all_persons']}
    assert roles.get('郑一') == 'bid_contact', roles


def t_personnel_same_line_label_not_swallowed():
    # A name followed by the NEXT field label on the same line must not
    # swallow it ('联系人：王五 传真：…' once became the 'name' 王五传真).
    # Lazy name classes stop at the boundary; the field-label suffix
    # blacklist in _is_person_name is the backstop.
    p = m.extract_personnel('投标函\n联系人：王五 传真：(010)1234-5678\n')
    roles = {x['name']: x['role'] for x in p['all_persons']}
    assert roles.get('王五') == 'bid_contact', roles
    p2 = m.extract_personnel('授权委托书\n委托代理人：李四 电话：13800000000\n')
    assert p2['authorized_rep'] == '李四', p2


def t_price_wei_shi_connectors():
    r = m.extract_prices('开标一览表\n投标总价为98万元\n')
    assert r['totalPriceInTax'] == 980000, r
    r2 = m.extract_prices('报价表\n总报价是350万\n')
    assert r2['totalPriceInTax'] == 3500000, r2


def t_price_traditional_numerals():
    # 萬/貳 traditional glyphs normalize before CN parsing.
    r = m.extract_prices('开标一览表\n大写：壹佰貳拾萬元整\n')
    assert r['totalPriceInTax'] == 1200000, r


def t_price_discount_rate_labels():
    r = m.extract_prices('报价函\n优惠率：5%\n')
    assert r['bidRate'] == '5%', r
    r2 = m.extract_prices('报价函\n折扣率：8.5%\n')
    assert r2['bidRate'] == '8.5%', r2
    # 个百分点 wording without a % sign.
    r3 = m.extract_prices('报价函\n下浮率：8.5个百分点\n')
    assert r3['bidRate'] == '8.5%', r3


def t_price_zandingjia_excluded():
    r = m.extract_prices('开标一览表\n暂定价：500000元\n投标总价：980000元\n')
    assert r['totalPriceInTax'] == 980000, r


def t_price_pretax_posttax_label_family():
    r = m.extract_prices('报价表\n税前价：1189450.24\n')
    assert r['totalPrice'] == 1189450.24, r
    r2 = m.extract_prices('报价表\n税后金额：1261819.76元\n')
    assert r2['totalPriceInTax'] == 1261819.76, r2


def t_price_xiaoxie_thin_space():
    r = m.extract_prices('开标一览表\n小写：1 261 819.76元\n')
    assert abs((r['totalPriceInTax'] or 0) - 1261819.76) < 0.01, r


# ══ 损坏 PDF 处理：MuPDF 诊断静音 + pypdf 打不开时降级 ══

def _make_classic_pdf_in(td, name):
    """Build the classic-xref fixture inside `td` and return (path, offsets).

    The basename check keeps the fixture write confined to the caller's
    temp dir (path-traversal guard)."""
    import os as _os
    assert _os.path.basename(name) == name
    path = _os.path.join(td, name)
    offsets = _make_classic_pdf(path)
    return path, offsets


def _corrupt_xref_offset(path, obj_offset):
    """Point one xref entry (the content-stream object) at a bogus offset —
    the 'cannot find object in xref' repair scenario."""
    raw = open(path, 'rb').read()
    raw = raw.replace(b'%010d 00000 n' % obj_offset, b'0099999999 00000 n')
    open(path, 'wb').write(raw)


def _destroy_startxref(path):
    """Break the startxref keyword so pypdf cannot open the file at all."""
    raw = open(path, 'rb').read().replace(b'startxref', b'startxreaF')
    open(path, 'wb').write(raw)


def _draw_cjk_grid_table(path):
    """Draw a 3x4 line grid with CJK cell text (china-s built-in font)."""
    import pymupdf
    d = pymupdf.open()
    p = d.new_page()
    rows, cols = 3, 4
    x0, y0, cw, ch = 72, 72, 90, 22
    for r in range(rows + 1):
        p.draw_line((x0, y0 + r * ch), (x0 + cols * cw, y0 + r * ch))
    for c in range(cols + 1):
        p.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + rows * ch))
    data = [['姓名', '职务', '联系电话', '备注'],
            ['王强', '项目经理', '13900000001', ''],
            ['李芳', '监理工程师', '13900000002', '']]
    for r, row in enumerate(data):
        for c, v in enumerate(row):
            if v:
                p.insert_text((x0 + c * cw + 4, y0 + r * ch + 15), v,
                              fontname='china-s')
    d.save(path)
    d.close()


def _draw_cjk_borderless_table(path):
    """Same 3x4 CJK table as _draw_cjk_grid_table but with no ruling lines —
    the case only the layout analyzer can see."""
    import pymupdf
    d = pymupdf.open()
    p = d.new_page()
    x0, y0, cw, ch = 72, 72, 90, 22
    data = [['姓名', '职务', '联系电话', '备注'],
            ['王强', '项目经理', '13900000001', ''],
            ['李芳', '监理工程师', '13900000002', '']]
    for r, row in enumerate(data):
        for c, v in enumerate(row):
            if v:
                p.insert_text((x0 + c * cw + 4, y0 + r * ch + 15), v,
                              fontname='china-s')
    d.save(path)
    d.close()


def _make_classic_pdf(path):
    """Hand-assemble a minimal classic-xref PDF (pymupdf writes xref streams,
    so the plain-table layout needed for these corruption tests is built
    byte-by-byte with correct offsets)."""
    objs = {
        1: b'<< /Type /Catalog /Pages 2 0 R >>',
        2: b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        3: (b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] '
            b'/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>'),
        5: b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    }
    stream = b'BT /F1 12 Tf 72 720 Td (bid price 980000 yuan test) Tj ET'
    objs[4] = b'<< /Length %d >>\nstream\n' % len(stream) + stream + b'\nendstream'
    out = bytearray(b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n')
    offsets = {}
    for n in sorted(objs):
        offsets[n] = len(out)
        out += b'%d 0 obj\n' % n + objs[n] + b'\nendobj\n'
    xref_off = len(out)
    n_objs = max(objs) + 1
    out += b'xref\n0 %d\n' % n_objs + b'0000000000 65535 f \n'
    for n in range(1, n_objs):
        out += b'%010d 00000 n \n' % offsets[n]
    out += (b'trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n'
            % (n_objs, xref_off))
    open(path, 'wb').write(bytes(out))
    return offsets


def t_pdf_broken_xref_still_extracts():
    # Damaged xref entry (bogus object offset): MuPDF repairs and extracts;
    # the 'MuPDF error: cannot find object in xref' console spam is silenced.
    import os as _os
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        path, offsets = _make_classic_pdf_in(td, 'broken.pdf')
        _corrupt_xref_offset(path, offsets[4])
        t = m.extract_text_with_tables(path)
        assert 'bid price' in t, repr(t[:80])


def t_pdf_pypdf_dead_mupdf_fallback():
    # startxref destroyed: pypdf raises, extraction degrades to MuPDF's
    # repair mode instead of failing the whole file.
    import os as _os
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        path, _ = _make_classic_pdf_in(td, 'dead.pdf')
        _destroy_startxref(path)
        t = m.extract_text_with_tables(path)
        assert 'bid price' in t, repr(t[:80])


def t_pdf_table_channel_union():
    # Structured table channel: find_tables(union=True) fuses the layout
    # analyzer's grids (pymupdf-layout) with line-based candidates. Draws a
    # real line grid with CJK cell text (china-s) and checks the pipe rows.
    # NOTE: the page-level text is read through pypdf in the full pipeline,
    # and pypdf cannot decode china-s without a ToUnicode map — so this test
    # exercises _fitz_page_tables_as_pipes directly.
    try:
        import pymupdf
    except ImportError:
        return  # pymupdf absent: channel disabled in production too
    import os as _os
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        path, _ = _make_classic_pdf_in(td, 'table.pdf')
        _draw_cjk_grid_table(path)
        d2 = pymupdf.open(path)
        pipes = m._fitz_page_tables_as_pipes(d2[0])
        assert '姓名' in pipes and '王强' in pipes and '项目经理' in pipes, pipes


def t_pdf_table_layout_modes():
    # PDF_TABLE_LAYOUT gating. The layout analyzer costs ~5-13x the line finder
    # and, once consulted, suppresses line-only tables: its classifier returns
    # no 'table' box for diagram-ish pages, and the default use_layout=True
    # path then bails out with an empty TableFinder, discarding rows the line
    # finder had already recovered. So 'auto' has to be a superset of 'off' —
    # escalate for borderless tables, never trade away a line-found one.
    try:
        import pymupdf
    except ImportError:
        return  # pymupdf absent: channel disabled in production too
    # Activate the layout analyzer exactly as _open_fitz() does. This test
    # calls _fitz_page_tables_as_pipes directly, so unlike a real extraction
    # nothing has imported pymupdf.layout for it — without this the borderless
    # case would fail for a reason that has nothing to do with the gating.
    try:
        import pymupdf.layout  # noqa: F401
    except Exception:
        pass  # optional: union degrades to line-based
    import tempfile
    saved = m.PDF_TABLE_LAYOUT
    try:
        with tempfile.TemporaryDirectory() as td:
            # Borderless: the line finder is blind, so 'off' yields nothing
            # and 'auto' must recover the table by escalating to the model.
            path = os.path.join(td, 'nolines.pdf')
            _draw_cjk_borderless_table(path)
            doc = pymupdf.open(path)
            m.PDF_TABLE_LAYOUT = 'off'
            assert m._fitz_page_tables_as_pipes(doc[0]) == '', '线框检测不应认出无框线表'
            m.PDF_TABLE_LAYOUT = 'auto'
            pipes = m._fitz_page_tables_as_pipes(doc[0])
            assert '王强' in pipes and '项目经理' in pipes, pipes
            doc.close()

            # Framed: the line pass answers, so 'auto' must return exactly its
            # rows and never escalate — that is the whole suppression guard.
            path2 = os.path.join(td, 'lines.pdf')
            _draw_cjk_grid_table(path2)
            doc2 = pymupdf.open(path2)
            m.PDF_TABLE_LAYOUT = 'off'
            plain = m._fitz_page_tables_as_pipes(doc2[0])
            assert '王强' in plain, plain
            m.PDF_TABLE_LAYOUT = 'auto'
            assert m._fitz_page_tables_as_pipes(doc2[0]) == plain
            doc2.close()
    finally:
        m.PDF_TABLE_LAYOUT = saved


# ══ 提取缓存与并行提取 ══

def t_extract_cache_roundtrip():
    # Extraction output is cached per content hash: the same file replays,
    # an edited file does not, and the page budget is part of the key.
    import tempfile
    saved_dir, saved_on = m.EXTRACT_CACHE_DIR, m.EXTRACT_CACHE_ENABLED
    try:
        with tempfile.TemporaryDirectory() as td:
            m.EXTRACT_CACHE_DIR = td
            m.EXTRACT_CACHE_ENABLED = True
            p = os.path.join(td, 'sample.txt')
            with open(p, 'w', encoding='utf-8') as f:
                f.write('投标人名称：测试科技有限公司\n法定代表人：张三\n')
            first = m.extract_text_with_tables(p)
            assert '张三' in first, first

            # Count real parses so a "hit" is proven, not merely consistent.
            parses = []
            real = m._extract_text_uncached

            def _counting(*a, **k):
                parses.append(1)
                return real(*a, **k)

            m._extract_text_uncached = _counting
            try:
                # Same bytes at a *different* path still hits: the key is the
                # content, not the path — uploads are re-saved under a fresh
                # random filename on every analysis, so a path or mtime key
                # would never hit in practice.
                p2 = os.path.join(td, 'copy.txt')
                with open(p2, 'w', encoding='utf-8') as f:
                    f.write('投标人名称：测试科技有限公司\n法定代表人：张三\n')
                assert m.extract_text_with_tables(p2) == first
                assert not parses, '同内容不同路径应命中缓存'

                # Changed bytes → different key → real extraction.
                with open(p, 'w', encoding='utf-8') as f:
                    f.write('投标人名称：另一家公司\n')
                changed = m.extract_text_with_tables(p)
                assert len(parses) == 1, '内容已变却未重新提取'
                assert '另一家公司' in changed, changed
            finally:
                m._extract_text_uncached = real

            # Different settings are a different key for identical bytes.
            assert (m._extract_cache_key(p, 10)
                    != m._extract_cache_key(p, 20)), 'max_pages 未纳入缓存键'
    finally:
        m.EXTRACT_CACHE_DIR, m.EXTRACT_CACHE_ENABLED = saved_dir, saved_on


def t_extract_cache_ignores_empty():
    # Empty results are not cached — that is also what a transient failure
    # looks like, and re-extracting a document that yielded nothing is cheap.
    import tempfile
    saved_dir, saved_on = m.EXTRACT_CACHE_DIR, m.EXTRACT_CACHE_ENABLED
    try:
        with tempfile.TemporaryDirectory() as td:
            m.EXTRACT_CACHE_DIR = td
            m.EXTRACT_CACHE_ENABLED = True
            p = os.path.join(td, 'empty.txt')
            open(p, 'w', encoding='utf-8').close()
            assert m.extract_text_with_tables(p) == ''
            assert not [n for n in os.listdir(td) if n.endswith('.txt')
                        and n != 'empty.txt'], '空结果不应写入缓存'
    finally:
        m.EXTRACT_CACHE_DIR, m.EXTRACT_CACHE_ENABLED = saved_dir, saved_on


def t_extract_many_matches_sequential():
    # _extract_many is the parallel path. It must return texts aligned with
    # its input despite workers finishing out of order, and an unreadable
    # file must yield '' instead of raising — the sequential loop it replaces
    # swallowed per-file errors the same way.
    try:
        import app  # noqa: F401 — the spawn path re-imports this module
    except Exception:
        return
    import tempfile
    saved_dir, saved_on = m.EXTRACT_CACHE_DIR, m.EXTRACT_CACHE_ENABLED
    try:
        with tempfile.TemporaryDirectory() as td:
            m.EXTRACT_CACHE_DIR = td
            m.EXTRACT_CACHE_ENABLED = False   # measure the path, not the cache
            paths = []
            for i, name in enumerate(('甲公司', '乙公司', '丙公司')):
                p = os.path.join(td, f'{i}.txt')
                with open(p, 'w', encoding='utf-8') as f:
                    f.write(f'投标人名称：{name}\n法定代表人：负责人{i}\n'
                            * (i + 1))
                paths.append(p)
            paths.append(os.path.join(td, 'missing.txt'))  # unreadable

            def _sequential_reference(ps):
                # Mirrors the loop _extract_many replaced: a per-file failure
                # contributed nothing rather than propagating.
                out = []
                for p in ps:
                    try:
                        out.append(m.extract_text_with_tables(p) or '')
                    except Exception:
                        out.append('')
                return out

            expected = _sequential_reference(paths)
            got = m._extract_many(paths)
            assert got == expected, [len(g) for g in got]

            seen = []
            m._extract_many(paths, on_file_done=lambda i, t: seen.append(i))
            assert sorted(seen) == list(range(len(paths))), seen
    finally:
        m.EXTRACT_CACHE_DIR, m.EXTRACT_CACHE_ENABLED = saved_dir, saved_on


def t_reference_filter_normalized():
    # _is_in_reference takes *pre-normalized* reference documents (it used to
    # re-normalize them on every call, which profiling put at 94% of the whole
    # analysis). The comparison must still see through whitespace, full-width
    # punctuation and case on both sides.
    ref = ['本项目采用三层架构设计，核心交换节点采用双机热备冗余部署策略。']
    ref_norm = [m._normalize_for_match(r) for r in ref]

    same_content = '本项目采用三层架构设计， 核心交换节点采用双机热备冗余部署策略。'
    assert m._is_in_reference(same_content, ref_norm), '空白差异应仍判为同一段'
    assert not m._is_in_reference('完全无关的另一段技术描述，用于确认不会误判', ref_norm)
    assert not m._is_in_reference('三层架构', ref_norm), '过短不应判为命中'
    assert not m._is_in_reference(same_content, []), '无参照文件时不应命中'


def t_extract_worker_sizing():
    # Worker count must adapt to the machine it lands on: bounded by the batch
    # size, the usable CPU count, the memory budget and the hard cap — and
    # never 0, never more than the batch. The memory term scales with input
    # size because a worker's footprint does (measured 458MB for a 2MB PDF,
    # 903MB for a 242MB one).
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        tiny = []
        for i in range(6):
            p = os.path.join(td, f'{i}.pdf')
            with open(p, 'wb') as f:
                f.write(b'x' * 1024)
            tiny.append(p)
        big = os.path.join(td, 'big.pdf')
        with open(big, 'wb') as f:          # sparse: 200MB, no real disk use
            f.seek(200 * 1024 * 1024 - 1)
            f.write(b'\0')

        saved_cpu, saved_mem = m._usable_cpu_count, m._memory_budget_mb
        try:
            # Roomy machine: bounded by the batch and the cap, not by cores.
            m._usable_cpu_count = lambda: 64
            m._memory_budget_mb = lambda: 131072
            assert m._extract_worker_count(tiny) == min(6, m.EXTRACT_MAX_WORKERS)

            # Tight memory binds well below the CPU count.
            m._usable_cpu_count = lambda: 8
            m._memory_budget_mb = lambda: 1024
            assert m._extract_worker_count(tiny) == 2, m._extract_worker_count(tiny)

            # The same budget buys fewer workers once the inputs are large —
            # this is what stops a 4GB desktop from swapping on 240MB bids.
            m._memory_budget_mb = lambda: 4096
            assert m._extract_worker_count([big] * 6) == 4, m._extract_worker_count([big] * 6)

            # Never below one worker, however little memory is reported.
            m._memory_budget_mb = lambda: 1
            assert m._extract_worker_count(tiny) == 1

            # Unknown memory (no API on the platform) → CPU decides.
            m._memory_budget_mb = lambda: None
            m._usable_cpu_count = lambda: 3
            assert m._extract_worker_count(tiny) == 3

            # A single file never deserves a pool.
            assert m._extract_worker_count([tiny[0]]) == 1
        finally:
            m._usable_cpu_count, m._memory_budget_mb = saved_cpu, saved_mem

    # On whatever machine this runs on, both probes must return something
    # usable rather than raising.
    assert m._usable_cpu_count() >= 1
    budget = m._memory_budget_mb()
    assert budget is None or budget > 0, budget


# ══ 围串标信号：电话池 / 环境噪声 / 新增判定条款 ══

def t_phone_pool_includes_landlines():
    # The cross-matching pool used to be mobile-only, so a shared 座机 on the
    # 供应商基本情况表 (shared contact) could never pair two documents.
    r = m.extract_personnel('联系人：陈二 电话：010-88886666 传真：010-77775555')
    assert '010-88886666' in r['phones'], r['phones']
    assert '010-77775555' in r['phones'], r['phones']
    # Account / ID / date digit runs must NOT reach the pool.
    r2 = m.extract_personnel('账号：11000000000000000999 身份证号：320101199001011234')
    assert not [p for p in r2['phones'] if len(p) > 13], r2['phones']


def t_env_demotion_prefers_reference_doc():
    # A value the TENDER document reprints is environmental. A value shared by
    # EVERY bidder but absent from the tender document is the strongest signal
    # there is — a shared contact is *defined* as the shared case, so the old frequency
    # rule (which degenerates to "in every document" at 3 bidders) deleted
    # exactly the evidence it existed to protect.
    shared = '13911112222'
    pools = {n: {'phones': [shared]} for n in ('甲', '乙', '丙')}
    ref = ['招标代理联系电话：01000000000']
    removed = m._demote_environmental_pool_values(pools, list(pools), ref_texts=ref)
    assert removed == set(), removed
    assert all(pools[n]['phones'] == [shared] for n in pools)

    # …but a value that IS in the tender document still goes.
    pools2 = {n: {'phones': ['01000000000', '13900000001']} for n in ('甲', '乙', '丙')}
    removed2 = m._demote_environmental_pool_values(pools2, list(pools2), ref_texts=ref)
    assert removed2 == {'01000000000'}, removed2
    assert all(pools2[n]['phones'] == ['13900000001'] for n in pools2)

    # No reference supplied → frequency fallback, unchanged from before.
    pools3 = {n: {'phones': [shared, f'1390000000{i}']} for i, n in enumerate('甲乙丙')}
    assert m._demote_environmental_pool_values(pools3, list(pools3)) == {shared}


def t_project_manager_parenthesized_hint():
    # '（项目经理姓名）吴九' — bracket hint then fill, in free-standing 承诺书
    # prose no section marker covers. 丙公司's 承诺书 names 吴九 while its
    # 简历表 names 郑明; only the 承诺书 ties the file to 乙公司 (标注项九).
    r = m.extract_personnel(
        '我方拟派往（工程名称）某改造工程工程 的项目经理 （项目经理姓名）吴九 '
        '现阶段没有担任任何在施建设工程项目的项目经理。')
    pms = [p['name'] for p in r['all_persons'] if p['role'] == 'project_manager']
    assert pms == ['吴九'], pms


def t_non_name_business_nouns_rejected():
    # Table-cell fragments that survived the surname + length checks used to
    # reach all_persons — '经营范围' even cross-matched across all three
    # bidders as a team_member.
    for bad in ('经营范围', '管理模块', '通知用户', '成影响的'):
        assert not m._is_person_name(bad), bad
    for good in ('陈二', '张三', '李四', '蒋五', '周八', '冯强', '吴九'):
        assert m._is_person_name(good), good


def t_verdict_has_article34_and_mixing_clauses():
    # 条例第三十四条 (单位负责人为同一人) and 第四十条第五项 (文件混装) had no
    # clause at all: the 周八 legal_rep overlap was detected as a cross-match
    # yet appeared nowhere in the verdict.
    import tempfile
    saved = m.EXTRACT_CACHE_ENABLED
    try:
        m.EXTRACT_CACHE_ENABLED = False
        with tempfile.TemporaryDirectory() as td:
            paths = {}
            for name, tpl in (('甲', '姓名：周八 投标人名称：甲公司\n'),
                              ('乙', '姓名：周八 投标人名称：乙公司\n')):
                p = os.path.join(td, f'{name}.txt')
                with open(p, 'w', encoding='utf-8') as f:
                    f.write(tpl)
                paths[name] = p
            gt = {n: m.extract_text_with_tables(p) for n, p in paths.items()}
            res = m.run_full_analysis(list(paths.values()), [], group_map={n: [p] for n, p in paths.items()},
                                      group_texts=gt)
            clauses = {c['clause']: c for c in res['verdict']['clauses']}
            assert '第（三十四）条' in clauses, list(clauses)
            assert '第（五）项' in clauses, list(clauses)
    finally:
        m.EXTRACT_CACHE_ENABLED = saved


def t_total_price_ladder_detected():
    # 标注项六: three totals in equal steps. The pairwise loop only reported
    # EQUAL totals, so 295万/300万/305万 produced no finding at all.
    import tempfile
    saved = m.EXTRACT_CACHE_ENABLED
    try:
        m.EXTRACT_CACHE_ENABLED = False
        with tempfile.TemporaryDirectory() as td:
            paths = {}
            for name, row in (('甲', '01 | 甲公司 | 壹佰玖拾伍万元整 | 1950000 元'),
                              ('乙', '01 | 乙公司 | 贰百万元整 | 2000000 元'),
                              ('丙', '01 | 丙公司 | 贰佰零伍万元整 | 2050000 元')):
                p = os.path.join(td, f'{name}.txt')
                with open(p, 'w', encoding='utf-8') as f:
                    f.write(f'投标人名称：{name}公司\n7 报价一览表\n报价一览表\n'
                            f'包号 | 供应商名称 | 大写 | 小写\n{row}\n')
                paths[name] = p
            gt = {n: m.extract_text_with_tables(p) for n, p in paths.items()}
            for n, t in gt.items():
                got = m.extract_prices(t)['totalPriceInTax']
                assert got, f'{n} 未提取到总价'
            res = m.run_full_analysis(list(paths.values()), [],
                                      group_map={n: [p] for n, p in paths.items()},
                                      group_texts=gt)
            ev = next(c for c in res['verdict']['clauses']
                      if c['clause'] == '第（四）项-b')['evidence']
            assert any('等差数列' in e for e in ev), ev
    finally:
        m.EXTRACT_CACHE_ENABLED = saved


def t_docx_tables_keep_document_order():
    # Paragraphs used to be emitted first and every table appended after, so a
    # table was torn from the heading that introduced it (报价一览表 heading at
    # ~500, its price table pushed to ~16700).
    try:
        from docx import Document  # noqa: F401
    except ImportError:
        return
    import tempfile
    from docx import Document as _Doc
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, 'order.docx')
        d = _Doc()
        d.add_paragraph('7 报价一览表')
        tb = d.add_table(rows=1, cols=2)
        tb.rows[0].cells[0].text = '包号'
        tb.rows[0].cells[1].text = '2000000'
        d.add_paragraph('表后段落')
        d.save(p)
        t = m.extract_text_with_tables(p)
        assert t.index('报价一览表') < t.index('2000000') < t.index('表后段落'), repr(t)


def t_docx_image_ocr_is_selective():
    # Only .docx images whose NEIGHBOURING text marks them as evidence get
    # OCRed — a 商务标 carries ~128 images and OCRing all of them costs minutes
    # while most are seals and photos. The engine is stubbed: this test is
    # about the selection rule, not about OCR quality.
    try:
        import cv2  # noqa: F401
        from docx import Document  # noqa: F401
    except ImportError:
        return  # deps absent: image OCR is off in production too
    import tempfile
    import numpy as np
    from docx import Document as _Doc

    saved_ocr = m._get_ocr_image_fn
    saved_flag = m.DOCX_IMAGE_OCR
    saved_floor = m.DOCX_IMAGE_OCR_MIN_BYTES
    try:
        with tempfile.TemporaryDirectory() as td:
            img = os.path.join(td, 'i.png')
            cv2.imwrite(img, np.full((640, 480, 3), 200, np.uint8))
            ocr_calls = []

            def _fake_ocr(blob):
                ocr_calls.append(blob)
                return 'OCR_TEXTLINE'

            m.DOCX_IMAGE_OCR = True
            m.DOCX_IMAGE_OCR_MIN_BYTES = 0        # isolate the trigger rule
            m._get_ocr_image_fn = lambda: _fake_ocr

            p = os.path.join(td, 'img.docx')
            d = _Doc()
            d.add_paragraph('4 磋商保证金凭证/交款单据电子件')   # triggers
            d.add_picture(img)
            d.add_paragraph('施工组织设计概述')                  # does not
            d.add_picture(img)
            d.save(p)

            text = m.extract_text_with_tables(p)
            assert 'OCR_TEXTLINE' in text, text[-200:]
            assert len(ocr_calls) == 1, f'应只 OCR 触发词附近那张，实际 {len(ocr_calls)}'

            # Disabled entirely → no OCR call, no injected text.
            ocr_calls.clear()
            m.DOCX_IMAGE_OCR = False
            assert 'OCR_TEXTLINE' not in m.extract_text_with_tables(p)
            assert not ocr_calls
    finally:
        m._get_ocr_image_fn = saved_ocr
        m.DOCX_IMAGE_OCR = saved_flag
        m.DOCX_IMAGE_OCR_MIN_BYTES = saved_floor


def t_docx_image_ocr_skips_icons():
    # The byte floor keeps logos and rules out of the OCR queue.
    assert m.DOCX_IMAGE_OCR_MIN_BYTES > 0
    assert not m._DOCX_IMAGE_TRIGGER.search('施工组织设计')
    assert m._DOCX_IMAGE_TRIGGER.search('磋商保证金凭证/交款单据电子件')
    assert m._DOCX_IMAGE_TRIGGER.search('开户许可')


def t_docx_image_ocr_reports_progress():
    # The UI needs to know OCR is running — it is tens of seconds on a large
    # Word file, and without an event the progress bar simply sits there.
    try:
        import cv2  # noqa: F401
        from docx import Document  # noqa: F401
    except ImportError:
        return
    import tempfile
    import numpy as np
    from docx import Document as _Doc

    saved_ocr = m._get_ocr_image_fn
    saved_floor = m.DOCX_IMAGE_OCR_MIN_BYTES
    try:
        with tempfile.TemporaryDirectory() as td:
            img = os.path.join(td, 'i.png')
            cv2.imwrite(img, np.random.randint(0, 255, (640, 480, 3), dtype=np.uint8))
            m._get_ocr_image_fn = lambda: (lambda blob: 'OCR_TEXTLINE')
            m.DOCX_IMAGE_OCR_MIN_BYTES = 0
            p = os.path.join(td, 'p.docx')
            d = _Doc()
            d.add_paragraph('4 磋商保证金凭证/交款单据电子件')
            d.add_picture(img)
            d.save(p)

            events = []
            m.extract_text_with_tables(
                p, on_progress=lambda ph, cur, tot, ht, det: events.append(ph))
            assert 'docx_img_ocr_start' in events, events
            assert 'docx_img_ocr' in events, events
            assert events.index('docx_img_ocr_start') < events.index('docx_img_ocr'), events
    finally:
        m._get_ocr_image_fn = saved_ocr
        m.DOCX_IMAGE_OCR_MIN_BYTES = saved_floor


def t_parallel_progress_crosses_processes():
    # Workers cannot call the parent's callback, so events travel over a queue
    # handed to the pool at creation. Getting this wrong is silent: the run
    # succeeds and simply reports nothing (the first cut used terminate() in
    # its finally, which killed the workers before they flushed that queue).
    try:
        import cv2  # noqa: F401
        from docx import Document  # noqa: F401
    except ImportError:
        return
    import tempfile
    import numpy as np
    from docx import Document as _Doc
    saved_dir, saved_on = m.EXTRACT_CACHE_DIR, m.EXTRACT_CACHE_ENABLED
    try:
        with tempfile.TemporaryDirectory() as td:
            m.EXTRACT_CACHE_DIR = td
            m.EXTRACT_CACHE_ENABLED = False
            img = os.path.join(td, 'i.png')
            cv2.imwrite(img, np.random.randint(0, 255, (640, 480, 3), dtype=np.uint8))
            paths = []
            for n in range(2):
                p = os.path.join(td, f'{n}.docx')
                d = _Doc()
                d.add_paragraph('4 磋商保证金凭证/交款单据电子件')
                d.add_picture(img)
                d.save(p)
                paths.append(p)
            events = []
            m._extract_many(paths, on_progress=lambda *a: events.append(a[0]))
            assert 'docx_img_ocr_start' in events, events
    finally:
        m.EXTRACT_CACHE_DIR, m.EXTRACT_CACHE_ENABLED = saved_dir, saved_on


def t_display_name_strips_upload_prefix():
    # Progress labels must name the document the user picked, not the storage
    # name _safe_save invented ('46e8be14_乙公司.docx').
    assert m._display_name('/tmp/uploads/46e8be14_乙公司.docx') == '乙公司.docx'
    assert m._display_name('ref_a1b2c3d4_招标文件.pdf') == '招标文件.pdf'
    # A name that merely starts with hex-ish text must survive intact.
    assert m._display_name('fedcba98x_报告.docx') == 'fedcba98x_报告.docx'
    assert m._display_name('上海全大-投标文件.pdf') == '上海全大-投标文件.pdf'
    assert m._display_name('') == ''


def t_cancel_reaches_extraction():
    # _extract_many must HAND cancel_event down to extract_text_with_tables.
    # It was dropped when the parallel rewrite replaced the per-file loops at
    # the call sites: Stop still set the flag, but nothing inside the
    # extractor ever looked — a 300-page scan kept OCRing (minutes) while the
    # button sat on 正在停止.
    try:
        from docx import Document  # noqa: F401
    except ImportError:
        return
    import tempfile
    import threading
    import time
    from docx import Document as _Doc

    saved_dir, saved_on = m.EXTRACT_CACHE_DIR, m.EXTRACT_CACHE_ENABLED
    try:
        with tempfile.TemporaryDirectory() as td:
            m.EXTRACT_CACHE_DIR = td
            m.EXTRACT_CACHE_ENABLED = False
            paths = []
            for n in range(2):
                p = os.path.join(td, f'{n}.docx')
                d = _Doc()
                for i in range(60):
                    d.add_paragraph(f'第 {i} 段内容')
                d.save(p)
                paths.append(p)

            # Sequential (single file) and parallel (two files) must both
            # abort near-instantly once the flag is up.
            for label, batch in (('串行', paths[:1]), ('并行', paths)):
                ev = threading.Event()
                ev.set()
                t0 = time.time()
                try:
                    m._extract_many(batch, cancel_event=ev)
                    raise AssertionError(f'{label}: 取消未生效，正常返回了')
                except m.AnalysisCancelled:
                    pass
                assert time.time() - t0 < 5, f'{label}: 取消响应过慢'
    finally:
        m.EXTRACT_CACHE_DIR, m.EXTRACT_CACHE_ENABLED = saved_dir, saved_on


def t_docx_image_ocr_zero_means_unlimited():
    # 0 must mean "unlimited" here as it does for OCR_MAX_PAGES /
    # OCR_TIME_BUDGET / MAX_PDF_PAGES. A bare `if done >= MAX` turns 0 into
    # "OCR nothing" — the exact opposite of what anyone setting 0 in this
    # codebase expects.
    try:
        import cv2  # noqa: F401
        from docx import Document  # noqa: F401
    except ImportError:
        return
    import tempfile
    import numpy as np
    from docx import Document as _Doc

    saved_ocr = m._get_ocr_image_fn
    saved_floor = m.DOCX_IMAGE_OCR_MIN_BYTES
    saved_max = m.DOCX_IMAGE_OCR_MAX
    saved_budget = m.DOCX_IMAGE_OCR_BUDGET
    try:
        with tempfile.TemporaryDirectory() as td:
            m._get_ocr_image_fn = lambda: (lambda blob: 'OCR_TEXTLINE')
            m.DOCX_IMAGE_OCR_MIN_BYTES = 0
            m.DOCX_IMAGE_OCR_MAX = 0          # unlimited
            m.DOCX_IMAGE_OCR_BUDGET = 0       # unlimited
            p = os.path.join(td, 'z.docx')
            d = _Doc()
            d.add_paragraph('4 磋商保证金凭证/交款单据电子件')
            for n in range(3):
                img = os.path.join(td, f'{n}.png')
                # distinct bytes so the de-dup does not collapse them
                cv2.imwrite(img, np.random.randint(0, 255, (200, 200, 3), dtype=np.uint8))
                d.add_picture(img)
            d.save(p)
            text = m.extract_text_with_tables(p)
            assert text.count('OCR_TEXTLINE') == 3, text.count('OCR_TEXTLINE')
    finally:
        m._get_ocr_image_fn = saved_ocr
        m.DOCX_IMAGE_OCR_MIN_BYTES = saved_floor
        m.DOCX_IMAGE_OCR_MAX = saved_max
        m.DOCX_IMAGE_OCR_BUDGET = saved_budget


def t_report_appendix_covers_new_clauses():
    # The appendix is what a reviewer reads to check the verdict's legal
    # footing: it must carry 第三十四条 (the new clause) and must not still
    # claim 第（五）项 needs manual review — it is auto-detected now.
    assert any('第三十四条' in x or '单位负责人为同一人' in x
               for x in m._REGULATION_ARTICLE_34), m._REGULATION_ARTICLE_34
    assert any('相互混装' in x for x in m._REGULATION_ARTICLE_40)
    # Weights must cover every clause the verdict can emit.
    for clause in ('第（三十四）条', '第（一）项', '第（二）项', '第（三）项',
                   '第（五）项', '第（四）项-a', '第（四）项-b'):
        assert clause in m._REPORT_CLAUSE_ADVICE, clause


def t_app_version_consistent_across_release_surfaces():
    # APP_VERSION 是版本号唯一来源：前端 footer 标注、静态资源缓存参数、
    # Release 产物文件名后缀、安装器版本都由它派生。iss 与 star.spec 没有
    # CI 自动同步，靠这个测试在本地拦截漂移。
    ver = m.APP_VERSION
    assert re.fullmatch(r'\d+\.\d+\.\d+', ver), ver
    # 渲染路径：'/' 必须把版本号带进页面（footer 标注 + 缓存参数）
    resp = m.app.test_client().get('/')
    assert resp.status_code == 200, resp.status_code
    body = resp.data.decode('utf-8')
    assert f'v{ver}' in body, 'footer 版本标注缺失'
    assert f'style.css?v={ver}' in body and f'main.js?v={ver}' in body, body[-500:]
    # 打包侧：iss 安装器版本与 macOS Bundle 版本须与 APP_VERSION 一致
    root = os.path.dirname(os.path.abspath(m.__file__))
    with open(os.path.join(root, 'packaging', '星易查.iss'), encoding='utf-8') as f:
        assert f'#define MyAppVersion "{ver}"' in f.read(), '星易查.iss 版本未同步'
    with open(os.path.join(root, 'packaging', 'star.spec'), encoding='utf-8') as f:
        assert f"'CFBundleShortVersionString': '{ver}'" in f.read(), 'star.spec 版本未同步'


def main():
    # Tests must be hermetic: the extraction cache lives on disk between runs,
    # and a cached PDF extraction silently skips the very code path a test
    # exists to exercise (the damaged-PDF tests, for one, would pass without
    # ever reaching the repair path). The cache tests turn it back on for
    # themselves and restore this value afterwards.
    m.EXTRACT_CACHE_ENABLED = False
    for name, fn in sorted(globals().items()):
        if name.startswith('t_') and callable(fn):
            check(name, fn)
    print(f'\n{len(PASS)}/{len(PASS)} tests passed')


if __name__ == '__main__':
    main()
